"""Amend 'VRC network changes after anoxia-induced gasping.pptx' (e097): intro slides, current versions of the
population slides, an anoxia #2 section comparing it side by side with anoxia #1, and more slides per unit.

The existing slides are kept as they are except where a step below says otherwise. Final order:
  1   title (kept)
  2-3 intro: the question and hypothesis; four questions (transition_slides.INTRO); replaces the empty slide 2
  4   how to read: conditions timeline and epochs table, current names (replaces slide 3)
  5   population PETH (kept)
  6   K = 6 breath raster: current figure (old picture used the superseded epoch names)
  7   airflow (kept)
  8   epoch table: current names, native table (replaces the slide 7 image)
  9   gasp -> recovery timeline (kept)
  10-16 anoxia #2: section, both trials' timelines, then anoxia #1 vs #2 side by side: population PETH, K = 6 breath
      raster, airflow, epoch table, gasp -> recovery timeline
  17  units section (kept), then per unit: the kept raster and PETH slides, then anoxia #1 vs #2 rasters, anoxia #2
      PETH, transition rasters #1 vs #2, transition PETHs #1 and #2 (#422: anoxia #1 transitions only). The old slide
      33 (after #510's raster) showed #435's PETH and title; it now shows #510's.
Every new or changed slide was approved by the user from a rendered preview (2026-10-22).

  cibrrig python amend_vrc_deck.py <in.pptx> <out.pptx>     (python-pptx on PYTHONPATH)
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import transition_slides as ts  # noqa: E402

OUT = Path("C:/Users/lpascu/resp_analysis/outputs/e097")
A, B = OUT / "population_dynamics_v2", OUT / "population_dynamics_anoxia2"
N2_ON = {"A": 1800.0, "B": 2310.0}
INK, INK2, GRID = RGBColor(0x1F, 0x1F, 0x1F), RGBColor(0x52, 0x51, 0x4E), RGBColor(0xD9, 0xD8, 0xD4)
TRANSITION = {"pre-gasp": "#9b3b78", "post-gasp": "#2c8a96"}
COLOURS_B = {"recovery 3": "#1c9157", "recovery 4": "#3f8f5a", "recovery 5": "#0f6b3a"}
UNITS = [(331, "Inspiratory"), (362, "Inspiratory"), (459, "Inspiratory"), (312, "Expiratory"), (363, "Expiratory"),
         (374, "Expiratory"), (341, "Expiratory"), (350, "Expiratory"), (353, "Expiratory"), (391, "Expiratory"),
         (435, "Expiratory"), (510, "Expiratory"), (422, "GASPING ONLY"), (425, "Tonic"), (434, "Tonic")]
W = 13.333


def rgb(hexstr):
    return RGBColor.from_string(hexstr.lstrip("#").upper())


# ------------------------------------------------------------------------------------------- primitives

def title_only(prs, title, size=32, unit_style=False, compact=False):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    t = s.shapes.title
    if unit_style:                                   # as the existing unit slides: top-left, large
        t.left, t.top, t.width, t.height = Inches(0.11), Inches(-0.28), Inches(13.0), Inches(1.45)
    if compact:                                      # two-line message title above a full-height figure
        t.left, t.top, t.width, t.height = Inches(0.4), Inches(0.15), Inches(12.55), Inches(1.1)
    t.text = title
    for p in t.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
    return s


def text(slide, x, y, w, h, runs, size=14, colour=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         italic=False):
    """runs: a string, or a list of paragraphs, each a string or a list of (text, bold) pairs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, b in ([(para, bold)] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.italic = Pt(size), b, italic
            r.font.color.rgb = colour
    return tb


def rect(slide, x, y, w, h, colour, shape=MSO_SHAPE.RECTANGLE, line=None):
    r = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = rgb(colour) if isinstance(colour, str) else colour
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = line
    r.shadow.inherit = False
    return r


def picture(slide, path, x, y, w, h):
    """Fit the image inside the box (x, y, w, h), centred."""
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw),
                                    Inches(ph))


def side_by_side(slide, left, right, y=1.3, h=5.9, labels=("Anoxia #1", "Anoxia #2")):
    gap, x0 = 0.25, 0.3
    w = (W - 2 * x0 - gap) / 2
    for j, (path, lab) in enumerate(zip((left, right), labels)):
        x = x0 + j * (w + gap)
        text(slide, x, y, w, 0.35, lab, size=16, bold=True)
        rect(slide, x, y + 0.36, w, 0.03, GRID)
        if path is not None and Path(path).exists():
            picture(slide, path, x, y + 0.45, w, h - 0.45)
        else:
            text(slide, x, y + 1.2, w, 1, "no anoxia #2 partner (unit not matched)", size=14, colour=INK2)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def table(slide, x, y, w, h, header, rows, col_w=None, size=12, header_fill="#F2F2F0", group=None):
    """header: list of str; rows: list of lists; group: [(label, first col, last col)] for a top header row."""
    n_head = 2 if group else 1
    shape = slide.shapes.add_table(len(rows) + n_head, len(header), Inches(x), Inches(y), Inches(w), Inches(h))
    tb = shape.table
    tb.first_row = True
    if col_w:
        for j, cw in enumerate(col_w):
            tb.columns[j].width = Inches(cw)

    def put(cell, s, bold=False, fill=None, align=PP_ALIGN.LEFT):
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = s
        r.font.size, r.font.bold = Pt(size), bold
        r.font.color.rgb = INK
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill) if fill else RGBColor(0xFF, 0xFF, 0xFF)
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if group:
        for j in range(len(header)):
            put(tb.cell(0, j), "", fill=header_fill)
        for lab, a, b in group:
            tb.cell(0, a).merge(tb.cell(0, b))
            put(tb.cell(0, a), lab, bold=True, fill=header_fill, align=PP_ALIGN.CENTER)
    for j, hd in enumerate(header):
        put(tb.cell(n_head - 1, j), hd, bold=True, fill=header_fill, align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            put(tb.cell(i + n_head, j), str(v), align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER)
    return shape


# ------------------------------------------------------------------------------------------------- data

def blocks(run, colours):
    """Contiguous runs of breaths in one condition: (name, start raw s, end raw s, n breaths, breaths/min)."""
    T = pd.read_csv(run / "transitions" / "transition_breaths.tsv", sep="\t")
    T["condition"] = T["condition"].fillna("")
    c, on = T["condition"].to_numpy(), T["on_raw_s"].to_numpy()
    change = np.r_[0, np.flatnonzero(c[1:] != c[:-1]) + 1, c.size]
    out = []
    for a, b in zip(change[:-1], change[1:]):
        if not c[a]:
            continue
        end = on[b] if b < on.size else on[-1]
        out.append((c[a], on[a], end, int(b - a), 60 / np.median(T["ibi_s"].to_numpy()[a:b]), colours.get(c[a])))
    return out


def colours_for(run):
    tw_ctx = {**ts.tw.TRANSITION_COLOURS}
    import epoch_change_points as ecp
    return {**ecp.COLOURS, **COLOURS_B, **tw_ctx}


def timeline(slide, bl, n2, x0, y, w, t0, t1, label="", h=0.42, size=10, below=False):
    """Coloured blocks on a shared axis of seconds from the N2 switch; labels above, staggered."""
    xs = lambda t: x0 + w * (t - n2 - t0) / (t1 - t0)  # noqa: E731
    if label:
        text(slide, x0 - 1.45, y, 1.35, h, label, size=13, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
    last_end = -1e9
    level = 0
    for name, a, b, n, rate, col in bl:
        rect(slide, xs(a), y, max(xs(b) - xs(a), 0.02), h, col or "#FFFFFF")
        lx = xs(a)
        level = level + 1 if lx < last_end + 0.05 else 0
        ty = (y + h + 0.03 + 0.27 * (level % 3)) if below else (y - 0.3 - 0.27 * (level % 3))
        tb = text(slide, lx, ty, 2.2, 0.27, name, size=size, colour=INK)
        last_end = lx + 0.085 * size / 10 * len(name)
        del tb
    return xs


# ------------------------------------------------------------------------------------------------ slides

def intro_question(prs):
    I = ts.INTRO
    s = title_only(prs, I["title"], size=32)
    text(s, 0.92, 1.85, 6, 0.3, "THE QUESTION", size=12, bold=True, colour=INK2)
    text(s, 0.92, 2.15, 11.6, 0.8, I["question"], size=20)
    text(s, 0.92, 3.15, 6, 0.3, "CENTRAL HYPOTHESIS", size=12, bold=True, colour=INK2)
    bars = ["#6a6a66", "#d93636", "#77c06a"]
    for i, h in enumerate(I["hypothesis"]):
        y = 3.5 + i * 1.12
        rect(s, 0.92, y, 0.08, 0.95, bars[i])
        lead, rest = h.split(" ", 1)
        text(s, 1.15, y, 5.6, 0.95, [[(lead + " ", True), (rest, False)]], size=15, anchor=MSO_ANCHOR.MIDDLE)
    fig = A / "transitions" / "slides" / "intro_trajectories.png"
    pic = picture(s, fig, 7.0, 3.45, 5.9, 2.7)
    text(s, 7.0, pic.top / 914400 + pic.height / 914400 + 0.08, 5.9, 0.5, I["figure_caption"], size=11, colour=INK2,
         italic=True)
    notes(s, "Question and hypothesis paraphrased from the K00 Research Strategy (Innovation; Aims 1–3).")
    return s


WHERE = ["Population PETHs; unit slides through the transitions",
         "K = 6 breath raster; gasp → recovery timeline; unit slides",
         "Gasp → recovery timeline (heart rate row)",
         "Anoxia #2 section"]


def intro_subquestions(prs):
    I = ts.INTRO
    s = title_only(prs, I["subquestions_title"], size=32)
    n, top = len(I["subquestions"]), 1.95
    row = (7.05 - top) / n
    for i, q in enumerate(I["subquestions"]):
        y = top + i * row
        if i:
            rect(s, 0.92, y - 0.08, 11.6, 0.012, GRID)
        o = rect(s, 0.92, y + 0.05, 0.5, 0.5, "#52514E", shape=MSO_SHAPE.OVAL)
        o.text_frame.text = str(i + 1)
        p = o.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size, p.runs[0].font.bold = Pt(16), True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        text(s, 1.65, y, 7.9, 0.45, q["q"], size=18, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 1.65, y + 0.47, 7.9, row - 0.55, q["how"], size=13)
        text(s, 9.85, y + 0.05, 2.65, row - 0.15, [[(q["aim"], True)], [(WHERE[i], False)]], size=12, colour=INK2)
    notes(s, "Sub-questions mapped onto the K00 aims: network dynamics from eupnea to gasping (Aim 1), re-embedding "
             "and post-inspiratory recruitment during recovery (Aims 1–3), cardiorespiratory coupling (Aim 3); right "
             "column: where in this deck.")
    return s


def how_to_read(prs):
    s = title_only(prs, "How to read the unit slides: conditions and epochs (anoxia #1)", size=32)
    bl = blocks(A, colours_for(A))
    timeline(s, bl, N2_ON["A"], 0.92, 2.55, 11.5, -60, 360)
    text(s, 0.92, 3.02, 11.5, 0.3, "raw-run time 1740–2160 s; pre-gasp and post-gasp are the breaths between anoxia "
         "and the first gasp, and between the last gasp and early recovery (unit transition slides)", size=11,
         colour=INK2, italic=True)
    src = {"baseline": "60 s before the logged N2 switch", "anoxia": "15 s N2 + 10 s",
           "pre-gasp": "after anoxia, before the first gasp", "gasping": "Bush & Ramirez 2024 IBI rule",
           "post-gasp": "after the last gasp, before early recovery"}
    rows = [[n, f"{a:.0f}–{b:.0f} s", k, f"{r:.0f}", src.get(n, "change point, breathing + HMM state")]
            for n, a, b, k, r, _ in bl]
    table(s, 0.92, 3.5, 11.5, 0.33 * (len(rows) + 1), ["condition", "raw time", "breaths", "breaths/min",
                                                      "defined by"], rows, col_w=[2.2, 1.8, 1.2, 1.5, 4.8], size=12)
    notes(s, "Replaces the earlier image of this slide, which used the superseded epoch names (settled recovery, "
             "sigh stretch, active expiration).")
    return s


def metrics_rows(run, names):
    m = pd.read_csv(run / "figures" / "epoch_metrics.tsv", sep="\t", index_col=0)
    fmt = {"median firing rate vs baseline": "{:.2f}", "breathing modulation (median L)": "{:.2f}",
           "inspiratory-locked units": "{:.0f}", "correlation structure vs baseline (r)": "{:.2f}",
           "time in the recovery post-I state": "{:.0%}", "time in the recovery late-E state": "{:.0%}"}
    return m, [[lab] + [fmt[lab].format(m.loc[lab, n]) for n in names] for lab in fmt]


def epoch_table(prs):
    s = title_only(prs, "Lasting central changes do not support recovery of peripheral breathing metrics", size=32)
    names = ["baseline", "early recovery", "mid recovery", "slow breathing", "later recovery"]
    m, rows = metrics_rows(A, names)
    table(s, 1.19, 2.14, 10.95, 0.42 * 7, [""] + names, rows, col_w=[3.95] + [1.4] * 5, size=14)
    text(s, 1.19, 5.3, 10.95, 1.0,
         f"Median L and inspiratory units count every unit; a near-silent unit (one spike) reads as L = 1. With ≥ 20 "
         f"spikes per epoch: median L {m.loc['median L, units with >= 20 spikes', 'baseline']:.2f} at baseline, "
         + ", ".join(f"{m.loc['median L, units with >= 20 spikes', n]:.2f}" for n in names[1:])
         + "; inspiratory units " + ", ".join(f"{m.loc['inspiratory units, >= 20 spikes', n]:.0f}" for n in names)
         + ". States: K = 6 HMM, mean state probability.", size=11, colour=INK2, italic=True)
    notes(s, "Same values as the earlier image of this table; epochs renamed (early / mid recovery, slow breathing, "
             "later recovery) and a baseline column added. Source: population_dynamics_v2/figures/epoch_metrics.tsv.")
    return s


def section_b(prs):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Anoxia #2: the same neurons, a second anoxia"
    s.placeholders[1].text = ("27 of the 28 units matched by template; decoded with the anoxia #1 HMM and projected on "
                              "the anoxia #1 PCA axes")
    return s


def trials_timeline(prs):
    s = title_only(prs, "Anoxia #2 has 8 gasps (20 in anoxia #1) and recovers faster", size=32)
    xs = timeline(s, blocks(A, colours_for(A)), N2_ON["A"], 2.4, 2.45, 10.3, -60, 360, label="Anoxia #1")
    for t in range(-60, 361, 60):
        x = xs(t + N2_ON["A"])
        rect(s, x, 2.9, 0.01, 0.07, INK2)
        text(s, x - 0.4, 2.97, 0.8, 0.22, f"{t}", size=10, colour=INK2, align=PP_ALIGN.CENTER)
    text(s, 2.4, 3.2, 10.3, 0.25, "s from the N2 switch (anoxia #1 30:00 = 1800 s; anoxia #2 38:30 = 2310 s)",
         size=10, colour=INK2, align=PP_ALIGN.CENTER)
    timeline(s, blocks(B, colours_for(B)), N2_ON["B"], 2.4, 3.55, 10.3, -60, 360, label="Anoxia #2", below=True)
    text(s, 0.92, 5.05, 11.6, 1.8, [
        "Units: the anoxia #2 partners of the anoxia #1 singles (baseline templates, cosine ≥ 0.99), labelled by their "
        "anoxia #1 ids; #422 has no partner. The anoxia #2 sort has no unit review.",
        "HMM states are anoxia #1's (anoxia #2 decoded with the anoxia #1 K = 6 model); PCA on anoxia #1's axes.",
        "Anoxia #2's recovery split at change points as for anoxia #1; the stretches after early recovery are numbered "
        "(recovery 3: 113/min, the slowest)."], size=13)
    notes(s, "Anoxia #2 epochs (raw s): baseline 2250–2310, anoxia 2310–2335, pre-gasp 2335–2347, gasping 2348–2356 "
             "(8 gasps), post-gasp 2357–2371, early recovery 2371–2433, recovery 3 2433–2540, recovery 4 2540–2550, "
             "recovery 5 2550–2609.")
    return s


COMPARE = [
    ("Population PETHs: anoxia #2 starts in the recovery-type expiratory states, and its early recovery resembles "
     "anoxia #1's mid recovery",
     B / "figures" / "compare" / "population_peth_ion_compare.png",
     "Anoxia #2 baseline: 55 % of the time in the recovery late-E state (yellow) and 9 % in recovery post-I (blue); "
     "anoxia #1 baseline: 0 % in both. Anoxia #2 early recovery: firing 0.87 of baseline, post-I 26 %, late-E 49 %, "
     "close to anoxia #1's mid recovery (0.85, 32 %, 43 %); anoxia #1's early recovery was post-I-dominated (55 %)."),
    ("Breath by breath: anoxia #2's baseline is already in the recovery late-E state, and returns to it right after "
     "gasping", "figures/breath_raster_K6_pm400ms.png",
     "K = 6 Viterbi state on every breath, ±400 ms around inspiration onset (left) and offset (right). Anoxia #2 "
     "switches to the baseline late-E state (green) only around the end of N2 (~2325 s), before the transition."),
    ("Expiratory flow before inspiration recurs in anoxia #2's slowest recovery stretch (36 % of breaths; 63 % in "
     "anoxia #1's slow breathing)", "figures/airflow_by_condition.png",
     "Both panels made the same way (airflow_by_condition.png; the existing airflow slide uses the earlier "
     "airflow_by_period figure). Anoxia #2 recovery 3 (2433–2540 s) breathes at 113/min vs 158 at its baseline, as "
     "anoxia #1's slow breathing (106 vs 190)."),
    None,                                                           # the table slide
    ("Out of gasping: anoxia #1 keeps gasp-mode expiration for 11 breaths, anoxia #2 for 1",
     B / "figures" / "compare" / "timeline_gasp_to_recovery_compare.png",
     "Anoxia #1: expiration switches to the recovery post-I state at 1870.3 s, 8.5 s after the last gasp; anoxia #2 "
     "at 2358.1 s, 1.7 s after. The heart rate steps up ~9 s after the switch in both."),
]


def compare_slide(prs, title, rel, note):
    s = title_only(prs, title, size=26, compact=True)
    if isinstance(rel, Path):                       # one figure with both trials
        picture(s, rel, 0.3, 1.3, W - 0.6, 6.05)
    else:
        side_by_side(s, A / rel, B / rel, y=1.3, h=6.05)
    notes(s, note)
    return s


def compare_table(prs):
    s = title_only(prs, "Anoxia #2 recovers without anoxia #1's low-rate, strongly modulated early stage", size=28,
                   compact=True)
    na = ["baseline", "early recovery", "mid recovery", "slow breathing", "later recovery"]
    nb = ["baseline", "early recovery", "recovery 3", "recovery 4", "recovery 5"]
    _, ra = metrics_rows(A, na)
    _, rb = metrics_rows(B, nb)
    rows = [a + b[1:] for a, b in zip(ra, rb)]
    hdr = [""] + ["baseline", "early", "mid", "slow", "later"] + ["baseline", "early", "rec. 3", "rec. 4", "rec. 5"]
    table(s, 0.4, 1.6, 12.55, 0.42 * 8, hdr, rows, col_w=[3.05] + [0.95] * 10, size=12,
          group=[("Anoxia #1", 1, 5), ("Anoxia #2", 6, 10)])
    text(s, 0.4, 5.25, 12.5, 1.2, "Columns are recovery epochs in time order (anoxia #1: early, mid recovery, slow "
         "breathing, later recovery; anoxia #2: early recovery, recoveries 3–5). Anoxia #2 has 27 units (no #422); "
         "states are the anoxia #1 K = 6 states (mean probability). Median L counts every unit (see the anoxia #1 "
         "table note).", size=11, colour=INK2, italic=True)
    notes(s, "Source: epoch_metrics.tsv in population_dynamics_v2/figures and population_dynamics_anoxia2/figures "
             "(sorting_v2/trial_compare.py).")
    return s


def unit_slides(prs, u, kind):
    name = f"{kind} Unit #{u}"
    has_b = (B / "figures" / "units_slides" / f"unit_{u}_raster.png").exists()
    out = []
    if has_b:
        s = title_only(prs, f"{name} · anoxia #1 vs #2", size=36, unit_style=True)
        side_by_side(s, A / "figures" / "units_slides" / f"unit_{u}_raster.png",
                     B / "figures" / "units_slides" / f"unit_{u}_raster.png", y=1.0, h=6.35)
        notes(s, "Spikes on every breath, whole windows: anoxia #1 1740–2160 s, anoxia #2 2250–2610 s (its matched "
                 "partner, labelled by the anoxia #1 id).")
        out.append(s)
        s = title_only(prs, f"{name} · anoxia #2", size=36, unit_style=True)
        picture(s, B / "figures" / "units_slides" / f"unit_{u}_peth.png", 0.42, 1.17, 12.5, 5.75)
        notes(s, "Anoxia #2 PETHs and phase curves, conditions in time order (recovery 3–5 are anoxia #2's later "
                 "recovery stretches).")
        out.append(s)
    s = title_only(prs, f"{name} · through the transitions" + (", anoxia #1 vs #2" if has_b else ""), size=36,
                   unit_style=True)
    if has_b:
        side_by_side(s, A / "transitions" / "units_slides" / f"unit_{u}_raster.png",
                     B / "transitions" / "units_slides" / f"unit_{u}_raster.png", y=1.0, h=6.35)
    else:
        picture(s, A / "transitions" / "units_slides" / f"unit_{u}_raster.png", 0.32, 0.9, 8.7, 6.6)
        text(s, 9.4, 2.1, 3.7, 2, "No anoxia #2 partner: this unit wasn't matched in the anoxia #2 sort.", size=18)
    notes(s, "Each row one breath from late anoxia to early recovery, from 0.3 s before inspiration onset to the next "
             "onset.")
    out.append(s)
    for trial, run in (("anoxia #1", A), ("anoxia #2", B)):
        p = run / "transitions" / "units_slides" / f"unit_{u}_peth.png"
        if not p.exists() or (trial == "anoxia #2" and not has_b):
            continue
        s = title_only(prs, f"{name} · transitions, {trial}", size=36, unit_style=True)
        picture(s, p, 0.42, 1.17, 12.5, 5.75)
        notes(s, f"{trial}: PETHs and phase curves with pre-gasp and post-gasp as their own columns.")
        out.append(s)
    return out


# --------------------------------------------------------------------------------------------- assemble

def swap_picture(slide, path):
    """Replace the slide's (single) picture with another image of the same aspect, same box."""
    pic = next(sh for sh in slide.shapes if sh.shape_type in (13, 14) and hasattr(sh, "image"))
    box = (pic.left, pic.top, pic.width, pic.height)
    pic._element.getparent().remove(pic._element)
    slide.shapes.add_picture(str(path), *box)


def build(src, dst, only=None):
    prs = Presentation(src)
    old = list(prs.slides)
    new = {}
    new["intro1"], new["intro2"] = intro_question(prs), intro_subquestions(prs)
    new["read"] = how_to_read(prs)
    swap_picture(old[4], A / "figures" / "breath_raster_K6_pm400ms.png")
    # slide 33 followed #510's raster with #435's PETH and title: give it #510's
    swap_picture(old[32], A / "figures" / "units_slides" / "unit_510_peth.png")
    for sh in old[32].shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text = r.text.replace("#435", "#510")
    new["table"] = epoch_table(prs)
    new["b_section"], new["b_timeline"] = section_b(prs), trials_timeline(prs)
    for i, spec in enumerate(COMPARE):
        new[f"b_cmp{i}"] = compare_table(prs) if spec is None else compare_slide(prs, *spec)
    units = {u: unit_slides(prs, u, kind) for u, kind in UNITS}
    order = [old[0], new["intro1"], new["intro2"], new["read"], old[3], old[4], old[5], new["table"], old[7],
             new["b_section"], new["b_timeline"]] + [new[f"b_cmp{i}"] for i in range(len(COMPARE))] + [old[8]]
    for k, (u, _) in enumerate(UNITS):
        order += [old[9 + 2 * k], old[10 + 2 * k]] + units[u]
    if only is not None:
        order = [order[i - 1] for i in only]
    lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(e.rId): e for e in lst}
    keep = [by_part[s.part] for s in order]
    for e in list(lst):
        if e not in keep:
            prs.part.drop_rel(e.rId)
        lst.remove(e)
    for e in keep:
        lst.append(e)
    prs.save(dst)
    return len(keep)


if __name__ == "__main__":
    print(build(sys.argv[1], sys.argv[2]), "slides")
