"""Build (or refresh) an Excel workbook for tracking your inspection of every unit.

    python make_inspection_workbook.py C:/Users/lpascu/resp_analysis/outputs/e097/intact_window_v2

Writes <output_dir>/unit_inspection.xlsx with one row per unit from every sorter: label, whether
the other sorter found it, metrics, the bombcell criteria it fails, per-spike flag counts, your
verdicts from review_units.py, and yellow columns for your own inspection record. Every column has
a filter/sort dropdown (it is an Excel table).

Re-running refreshes the computed columns and keeps everything you typed in the yellow columns
(matched by sorter + unit id). The previous file is kept as unit_inspection_<time>.bak.xlsx; rows
for units that no longer exist move to the "Removed units" sheet. Close the workbook in Excel first.
Phy edits are not read back here (or by the pipeline or the review browser).
"""
from __future__ import annotations

import argparse
import shutil
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))

import spikeinterface.core as si  # noqa: E402

from npxsort.review import explain_bombcell, reason_summary  # noqa: E402

FONT = "Arial"
VERDICTS = ["Single unit", "Multi-unit", "Noise / artifact", "Split (part of another unit)", "Merged (>1 neuron)", "Unsure"]
YES_NO = ["Yes", "No"]
KEEP = ["Yes", "No", "Maybe"]
PHY_ACTIONS = ["None", "Split", "Merge", "Relabel", "Other"]
VIEWER_VERDICTS = {"single": "Single unit", "multi": "Multi-unit", "noise": "Noise / artifact",
                   "split": "Split (part of another unit)", "merged": "Merged (>1 neuron)", "unsure": "Unsure"}

# (header, width, number format, description) — computed columns, refreshed on every run
COMPUTED = [
    ("Sorter", 11, None, "kilosort4 or lupin"),
    ("Unit ID", 8, "0", "unit id in units.tsv, the analyzer and review_units.py"),
    ("Phy cluster ID", 9, "0", "cluster id shown in Phy (differs from Unit ID; Phy's si_unit_id column = Unit ID)"),
    ("Label", 10, None, "bombcell label: good / mua / non_soma / noise"),
    ("Shared with other sorter", 16, None, "Found by both (agreement >= 0.5 with the other sorter's unit), This sorter only, or Not compared (noise units are left out of the comparison)"),
    ("Matched unit (other sorter)", 11, "0", "the other sorter's unit paired with this one (one-to-one), if any"),
    ("Agreement", 10, "0.00", "matched spikes / (spikes in both units - matched spikes); spikes match within 0.4 ms"),
    ("Depth (um)", 9, "0", "estimated unit position along the probe (monopolar triangulation)"),
    ("Largest channel", 9, "0", "channel where the mean template is largest in uV (the viewer plots amplitudes here by default)"),
    ("Metrics channel", 9, "0", "the sorter's main channel, on which every metric and bombcell rule was computed"),
    ("Polarity", 10, None, "sign of the larger template deflection on the metrics channel"),
    ("Spikes", 8, "0", "spike count after cleaning"),
    ("Firing rate (Hz)", 9, "0.00", ""),
    ("SNR", 7, "0.0", "template peak / channel noise"),
    ("Amplitude median (uV)", 10, "0", "median spike amplitude on the metrics channel"),
    ("ISI violations ratio", 9, "0.000", "refractory violations relative to a Poisson process (1.5 ms)"),
    ("Sliding RP violation", 9, "0.000", "contamination at 90% confidence; blank = test inconclusive (counted as a fail)"),
    ("L-ratio", 8, "0.00", ""),
    ("Fails (bombcell)", 40, None, "the criteria that decided the label (empty for good)"),
    ("Footprint-mismatch spikes", 10, "0", "spikes whose amplitude pattern across channels doesn't match the template (review_units.py flag F); blank = not computed"),
    ("Duplicate spikes", 10, "0", "spikes also assigned to a nearby unit within 0.1 ms, beyond chance (flag X); blank = not computed"),
    ("Viewer verdict", 14, None, "your verdict in review_units.py (copied at generation time)"),
    ("Viewer note", 30, None, "your note in review_units.py (copied at generation time)"),
    ("Open in viewer", 9, None, "opens this unit in review_units.py (start it first; default port 8765)"),
]
# yellow columns you fill in; kept across regenerations
MANUAL = [
    ("Inspected", 10, None, YES_NO, "Yes once you've looked at the unit"),
    ("My verdict", 22, None, VERDICTS, "your call after inspection"),
    ("Keep for analysis", 10, None, KEEP, "whether to include the unit downstream"),
    ("Phy action", 10, None, PHY_ACTIONS, "what you changed in Phy, if anything (Phy edits are not read back automatically)"),
    ("Notes", 50, None, None, "free text"),
    ("Inspected on", 12, "yyyy-mm-dd", None, "date of inspection"),
]
YELLOW = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
HEADER_FILL = PatternFill("solid", start_color="1F4E78", end_color="1F4E78")
LABEL_FILLS = {"good": "C6EFCE", "mua": "FFEB9C", "non_soma": "E4DFEC", "noise": "D9D9D9"}


def unit_rows(out_dir: Path, sorter: str, other: str | None, scratch: Path | None) -> pd.DataFrame:
    folder = out_dir / sorter
    t = pd.read_csv(folder / "units.tsv", sep="\t", index_col="unit_id")
    a = si.load_sorting_analyzer(folder / "analyzer", load_extensions=False)
    ids = list(a.unit_ids)
    t = t.reindex(ids)
    dense = a.get_extension("templates").get_data(operator="average")
    ptp = dense.max(axis=1) - dense.min(axis=1)
    main = np.asarray(a.get_main_channels(outputs="index"), dtype=int)
    pol = ["positive" if dense[i, :, main[i]].max() > -dense[i, :, main[i]].min() else "negative" for i in range(len(ids))]

    thr_file = folder / "bombcell" / "thresholds_used.json"
    reasons = [""] * len(ids)
    if thr_file.exists():
        import json
        _, crit = explain_bombcell(t, json.loads(thr_file.read_text()))
        reasons = [reason_summary(crit[u], t.at[u, "label"]) for u in ids]

    phy = {}
    phy_map = folder / "phy" / "cluster_si_unit_ids.tsv"
    if phy_map.exists():
        m = pd.read_csv(phy_map, sep="\t")
        phy = dict(zip(m["si_unit_id"], m["cluster_id"]))

    n_mis = n_dup = None
    if scratch is not None and (scratch / "review_cache" / f"flags_{sorter}.npz").exists():
        z = np.load(scratch / "review_cache" / f"flags_{sorter}.npz", allow_pickle=False)
        unit_index = a.sorting.to_spike_vector()["unit_index"]
        if "rel" in z and z["rel"].size == unit_index.size:
            mis = (z["rel"] > 0.5) & (z["sig"] > 6.0)
            if "reversed" in z:
                mis |= z["reversed"]
            n_mis = np.bincount(unit_index[mis], minlength=len(ids))
            n_dup = np.bincount(unit_index[z["partner"] >= 0], minlength=len(ids))

    review = {}
    rf = out_dir / "review" / f"{sorter}_unit_review.tsv"
    if rf.exists():
        r = pd.read_csv(rf, sep="\t", dtype=str, keep_default_na=False)
        review = {int(x["unit_id"]): dict(x) for _, x in r.iterrows()}

    rows = []
    for i, u in enumerate(ids):
        r = t.loc[u]
        fbb = r.get("found_by_both")
        fbb = str(fbb).lower() == "true"
        shared = "Not compared (noise)" if r["label"] == "noise" else ("Found by both" if fbb else "This sorter only")
        match = r.get(f"match_{other}") if other else None
        rv = review.get(int(u), {})
        rows.append({
            "Sorter": sorter, "Unit ID": int(u), "Phy cluster ID": phy.get(int(u)),
            "Label": r["label"], "Shared with other sorter": shared,
            "Matched unit (other sorter)": None if pd.isna(match) else int(float(match)),
            "Agreement": None if pd.isna(r.get("agreement")) else float(r["agreement"]),
            "Depth (um)": float(r["depth_um"]), "Largest channel": int(ptp[i].argmax()), "Metrics channel": int(main[i]),
            "Polarity": pol[i], "Spikes": int(r["num_spikes"]), "Firing rate (Hz)": float(r["firing_rate"]),
            "SNR": _num(r.get("snr")), "Amplitude median (uV)": _num(r.get("amplitude_median")),
            "ISI violations ratio": _num(r.get("isi_violations_ratio")),
            "Sliding RP violation": _num(r.get("sliding_rp_violation")), "L-ratio": _num(r.get("l_ratio")),
            "Fails (bombcell)": reasons[i],
            "Footprint-mismatch spikes": None if n_mis is None else int(n_mis[i]),
            "Duplicate spikes": None if n_dup is None else int(n_dup[i]),
            "Viewer verdict": VIEWER_VERDICTS.get(rv.get("verdict", ""), rv.get("verdict", "")) if rv else "",
            "Viewer note": rv.get("note", "") if rv else "",
        })
    return pd.DataFrame(rows)


def _num(v):
    try:
        v = float(v)
    except (TypeError, ValueError):
        return None
    return None if np.isnan(v) else v


def read_manual(path: Path) -> tuple[dict, list[str]]:
    """Your yellow-column entries from an existing workbook, keyed by (sorter, unit id)."""
    if not path.exists():
        return {}, []
    wb = load_workbook(path)
    entries = {}
    for sheet in ("Units", "Removed units"):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
        col = {h: k for k, h in enumerate(header)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[col["Sorter"]] is None:
                continue
            vals = {h: row[col[h]] for h, *_ in MANUAL if h in col}
            if any(v not in (None, "") for v in vals.values()):
                entries[(str(row[col["Sorter"]]), int(row[col["Unit ID"]]))] = vals
    return entries, wb.sheetnames


def style_header(ws, ncols: int):
    for k in range(1, ncols + 1):
        c = ws.cell(row=1, column=k)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42


def build(out_dir: Path, port: int) -> Path:
    cfg = yaml.safe_load((out_dir / "config_used.yaml").read_text())
    scratch = Path(cfg["paths"]["scratch_dir"]) if cfg.get("paths", {}).get("scratch_dir") else None
    sorters = [n for n, s in cfg["sorters"].items() if s.get("enabled") and (out_dir / n / "units.tsv").exists()]
    frames = []
    for n in sorters:
        other = next((o for o in sorters if o != n), None)
        frames.append(unit_rows(out_dir, n, other, scratch))
    df = pd.concat(frames, ignore_index=True)

    path = out_dir / "unit_inspection.xlsx"
    manual, _ = read_manual(path)
    if path.exists():
        backup = out_dir / f"unit_inspection_{datetime.now():%Y%m%d_%H%M%S}.bak.xlsx"
        shutil.copy2(path, backup)
        print(f"Previous workbook kept as {backup.name}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Units"
    headers = [h for h, *_ in COMPUTED] + [h for h, *_ in MANUAL]
    ws.append(headers)
    n_comp = len(COMPUTED)
    for _, r in df.iterrows():
        key = (r["Sorter"], int(r["Unit ID"]))
        vals = [r.get(h) for h, *_ in COMPUTED[:-1]] + ["open"]
        vals = [None if (isinstance(v, float) and np.isnan(v)) else v for v in vals]
        mv = manual.pop(key, {})
        ws.append(vals + [mv.get(h) for h, *_ in MANUAL])
        row = ws.max_row
        link = ws.cell(row=row, column=n_comp)
        link.hyperlink = f"http://127.0.0.1:{port}/#sorter={r['Sorter']}&unit={int(r['Unit ID'])}"
        link.font = Font(name=FONT, size=10, color="0563C1", underline="single")
    last = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last):
        for k, c in enumerate(row):
            if k == n_comp - 1:
                continue
            c.font = Font(name=FONT, size=10)
            if k >= n_comp:
                c.fill = YELLOW
    for k, (h, width, fmt, *_) in enumerate(COMPUTED + [(m[0], m[1], m[2]) for m in MANUAL], start=1):
        ws.column_dimensions[get_column_letter(k)].width = width
        if fmt:
            for r in range(2, last + 1):
                ws.cell(row=r, column=k).number_format = fmt
    style_header(ws, len(headers))
    for k in range(n_comp + 1, len(headers) + 1):
        ws.cell(row=1, column=k).fill = PatternFill("solid", start_color="BF8F00", end_color="BF8F00")
    ws.freeze_panes = "C2"
    ref = f"A1:{get_column_letter(len(headers))}{last}"
    tab = Table(displayName="Units", ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    label_col = get_column_letter(headers.index("Label") + 1)
    for lab, color in LABEL_FILLS.items():
        ws.conditional_formatting.add(f"{label_col}2:{label_col}{last}",
                                      CellIsRule(operator="equal", formula=[f'"{lab}"'], fill=PatternFill("solid", start_color=color, end_color=color)))

    # drop-down lists for the yellow columns
    lists = wb.create_sheet("Lists")
    for j, (h, _, _, options, _) in enumerate(MANUAL, start=1):
        if not options:
            continue
        lists.cell(row=1, column=j, value=h).font = Font(name=FONT, bold=True)
        for i, o in enumerate(options, start=2):
            lists.cell(row=i, column=j, value=o).font = Font(name=FONT)
        col = get_column_letter(headers.index(h) + 1)
        dv = DataValidation(type="list", formula1=f"=Lists!${get_column_letter(j)}$2:${get_column_letter(j)}${len(options) + 1}",
                            allow_blank=True, showErrorMessage=True, errorTitle="Pick from the list",
                            error=f"Choose one of: {', '.join(options)}")
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
    date_col = get_column_letter(headers.index("Inspected on") + 1)
    dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True,
                         showErrorMessage=True, error="Enter a date, e.g. 2026-10-19")
    ws.add_data_validation(dvd)
    dvd.add(f"{date_col}2:{date_col}{last}")
    lists.sheet_state = "hidden"

    # entries for units that no longer exist are kept, not lost
    if manual:
        rem = wb.create_sheet("Removed units")
        rem.append(["Sorter", "Unit ID"] + [h for h, *_ in MANUAL])
        for (s_, u_), vals in sorted(manual.items()):
            rem.append([s_, u_] + [vals.get(h) for h, *_ in MANUAL])
        style_header(rem, 2 + len(MANUAL))
        print(f"{len(manual)} rows of your entries refer to units that no longer exist; kept in 'Removed units'.")

    summary_sheet(wb, headers, last, sorters)
    readme_sheet(wb, out_dir)
    order = ["Read me", "Summary", "Units", "Removed units", "Lists"]
    wb._sheets = sorted(wb._sheets, key=lambda sh: order.index(sh.title))
    wb.active = wb.sheetnames.index("Units")
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


def summary_sheet(wb, headers: list[str], last: int, sorters: list[str]):
    """Counts as formulas over the Units sheet, so they update as you fill in the yellow columns."""
    s = wb.create_sheet("Summary")
    col = {h: get_column_letter(headers.index(h) + 1) for h in headers}
    rng = lambda h: f"Units!${col[h]}$2:${col[h]}${last}"  # noqa: E731
    s.append(["Summary (live formulas over the Units sheet)"])
    s.append([])
    s.append(["Count"] + sorters)
    rows = [("Units", None, None)]
    rows += [(f"Label: {lab}", "Label", lab) for lab in ("good", "mua", "non_soma", "noise")]
    rows += [(f"Shared: {v}", "Shared with other sorter", v) for v in ("Found by both", "This sorter only", "Not compared (noise)")]
    rows += [("Inspected: Yes", "Inspected", "Yes")]
    rows += [(f"My verdict: {v}", "My verdict", v) for v in VERDICTS]
    rows += [(f"Keep for analysis: {v}", "Keep for analysis", v) for v in KEEP]
    for title, h, v in rows:
        line = [title]
        for so in sorters:
            if h is None:
                line.append(f'=COUNTIFS({rng("Sorter")},"{so}")')
            else:
                line.append(f'=COUNTIFS({rng("Sorter")},"{so}",{rng(h)},"{v}")')
        s.append(line)
    insp = 3 + rows.index(("Inspected: Yes", "Inspected", "Yes")) + 1
    s.append(["Inspected (fraction of units)"] + [f"=IF({get_column_letter(k + 2)}4=0,0,{get_column_letter(k + 2)}{insp}/{get_column_letter(k + 2)}4)" for k in range(len(sorters))])
    frac_row = s.max_row
    s.append([])
    s.append(["Good units found by both sorters"] + [f'=COUNTIFS({rng("Sorter")},"{so}",{rng("Label")},"good",{rng("Shared with other sorter")},"Found by both")' for so in sorters])
    for row in s.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=10, bold=c.row in (1, 3))
    s.cell(row=1, column=1).font = Font(name=FONT, size=12, bold=True)
    for r in range(4, s.max_row + 1):
        for k in range(2, 2 + len(sorters)):
            s.cell(row=r, column=k).number_format = "0"
    for k in range(2, 2 + len(sorters)):
        s.cell(row=frac_row, column=k).number_format = "0%"
    s.column_dimensions["A"].width = 38
    for k in range(2, 2 + len(sorters)):
        s.column_dimensions[get_column_letter(k)].width = 14


def readme_sheet(wb, out_dir: Path):
    r = wb.create_sheet("Read me")
    lines = [
        ("Unit inspection workbook", True),
        (f"Generated {datetime.now():%Y-%m-%d %H:%M} from {out_dir}", False),
        ("", False),
        ("How to use", True),
        ("Units sheet: one row per unit from every sorter. Every column header has a filter/sort dropdown "
         "(e.g. 'Shared with other sorter' = Found by both, or Label = good).", False),
        ("Edit only the yellow columns (Inspected, My verdict, Keep for analysis, Phy action, Notes, Inspected on). "
         "Inspected, My verdict, Keep for analysis and Phy action are drop-down lists.", False),
        ("'Open in viewer' opens that unit in review_units.py; start the viewer first (default port 8765).", False),
        ("Regenerate with: python make_inspection_workbook.py <output_dir>. Computed columns refresh; your yellow "
         "entries are kept (matched by Sorter + Unit ID). Close the workbook in Excel first.", False),
        ("Phy edits (splits, merges, labels) are saved in the phy folder and are NOT read back by this workbook, "
         "the pipeline or the review browser. Record them in 'Phy action' / 'Notes'.", False),
        ("Viewer verdict / Viewer note are copied from review_units.py when the workbook is generated.", False),
        ("", False),
        ("Example of a filled-in row (yellow columns)", True),
        ("Inspected: Yes | My verdict: Merged (>1 neuron) | Keep for analysis: No | Phy action: Split | "
         "Notes: positive gasping unit (ch148) mixed in at 103-110 s; split in Phy | Inspected on: 2026-10-19", False),
        ("", False),
        ("Columns", True),
    ]
    for text, bold in lines:
        r.append([text])
        r.cell(row=r.max_row, column=1).font = Font(name=FONT, size=12 if text == lines[0][0] else 10, bold=bold)
    for h, _, _, desc in COMPUTED:
        r.append([h, desc])
    for h, _, _, _, desc in MANUAL:
        r.append([h, f"(yours) {desc}"])
    for row in r.iter_rows(min_row=len(lines) + 1):
        row[0].font = Font(name=FONT, size=10, bold=True)
        row[1].font = Font(name=FONT, size=10)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    for row in r.iter_rows(max_row=len(lines)):
        row[0].alignment = Alignment(wrap_text=False)
    r.column_dimensions["A"].width = 30
    r.column_dimensions["B"].width = 100


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("output_dir", help="the run's paths.output_dir")
    p.add_argument("--port", type=int, default=8765, help="review_units.py port used in the 'Open in viewer' links")
    args = p.parse_args()
    path = build(Path(args.output_dir), args.port)
    print(f"Wrote {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
